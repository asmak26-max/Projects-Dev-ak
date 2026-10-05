"""
STEP 1 — Workbook loading.

Loads the SAME uploaded bytes twice, for two different purposes:
  - data_only=False -> formula TEXT preserved. This is the copy we run
    detection on and eventually mutate/save (fixes must edit the actual
    formula, not a calculated value).
  - data_only=True  -> last-cached CALCULATED value. Used only to build
    a human-readable preview (a teacher wants to see grades/totals, not
    raw formula strings). Note: if the file was never opened & saved in
    real Excel, these cached values won't exist and will read as None --
    a real limitation of reading .xlsx without a calculation engine.
"""
import openpyxl
from io import BytesIO


def load_both(file_bytes):
    """file_bytes: raw bytes of the uploaded .xlsx. Returns (formula_wb, value_wb)."""
    formula_wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=False)
    value_wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    return formula_wb, value_wb


def get_all_sheets(wb):
    return [(name, wb[name]) for name in wb.sheetnames]


def workbook_summary(wb):
    summary = []
    for name, ws in get_all_sheets(wb):
        formula_cells = sum(
            1
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        )
        summary.append({
            "sheet": name, "dimensions": ws.dimensions,
            "rows": ws.max_row, "cols": ws.max_column, "formula_cells": formula_cells,
        })
    return summary
